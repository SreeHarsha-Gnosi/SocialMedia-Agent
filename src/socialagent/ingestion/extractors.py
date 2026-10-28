"""Turn uploaded files into `Attachment`s. One small function per format; registry-driven."""
from __future__ import annotations

import base64
import io
from pathlib import Path
from typing import Callable

import pandas as pd
from docx import Document

from socialagent.core.models import Attachment

MAX_TEXT_CHARS = 200_000
IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
               ".gif": "image/gif", ".webp": "image/webp"}


class UnsupportedFileError(ValueError):
    pass


def _clip(text: str) -> str:
    return text if len(text) <= MAX_TEXT_CHARS else text[:MAX_TEXT_CHARS] + "\n…[truncated]"


def _pdf(name: str, data: bytes) -> Attachment:
    # PDFs go to Claude natively so charts, tables and layout are understood.
    return Attachment(name, "pdf", media_type="application/pdf",
                      data_b64=base64.standard_b64encode(data).decode())


def _image(name: str, data: bytes) -> Attachment:
    return Attachment(name, "image", media_type=IMAGE_TYPES[Path(name).suffix.lower()],
                      data_b64=base64.standard_b64encode(data).decode())


def _docx(name: str, data: bytes) -> Attachment:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return Attachment(name, "text", text=_clip("\n".join(parts)))


def _describe_frame(df: pd.DataFrame) -> str:
    stats = df.describe(include="all").round(2).to_string() if not df.empty else "(empty)"
    return (f"Rows: {len(df)}, Columns: {list(df.columns)}\n\n"
            f"Summary statistics:\n{stats}\n\nFirst 200 rows:\n{df.head(200).to_csv(index=False)}")


def _csv(name: str, data: bytes) -> Attachment:
    return Attachment(name, "text", text=_clip(_describe_frame(pd.read_csv(io.BytesIO(data)))))


def _xlsx(name: str, data: bytes) -> Attachment:
    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
    body = "\n\n".join(f"## Sheet: {s}\n{_describe_frame(df)}" for s, df in sheets.items())
    return Attachment(name, "text", text=_clip(body))


def _text(name: str, data: bytes) -> Attachment:
    return Attachment(name, "text", text=_clip(data.decode("utf-8", errors="replace")))


_EXTRACTORS: dict[str, Callable[[str, bytes], Attachment]] = {
    ".pdf": _pdf, ".docx": _docx, ".csv": _csv, ".xlsx": _xlsx,
    ".txt": _text, ".md": _text, ".json": _text,
    **{ext: _image for ext in IMAGE_TYPES},
}

SUPPORTED_EXTENSIONS = sorted(e.lstrip(".") for e in _EXTRACTORS)


def extract(name: str, data: bytes) -> Attachment:
    fn = _EXTRACTORS.get(Path(name).suffix.lower())
    if fn is None:
        raise UnsupportedFileError(f"Unsupported file type: {name}")
    return fn(name, data)
